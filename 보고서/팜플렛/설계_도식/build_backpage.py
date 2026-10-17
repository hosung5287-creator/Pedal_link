# PedalLink 팜플렛에 뒷면(2페이지) "설계 한눈에"를 붙여 PedalLink_팜플렛_양면.pptx 를 만든다.
# 앞면(원본 PedalLink_팜플렛.pptx)은 건드리지 않는다. 글자는 모두 편집 가능한 텍스트 상자, 도식 3장은 PNG.
# 필요: python-pptx (pip install python-pptx), 같은 폴더의 architecture.png / flow.png / data.png
# 실행: python build_backpage.py
import io
import os

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'PedalLink_팜플렛.pptx')
OUT = os.path.join(HERE, '..', 'PedalLink_팜플렛_양면.pptx')
GITHUB = 'https://github.com/hosung5287-creator/Pedal_link'

INK, MUTED, LINE, BG2 = '0C1B27', '5B7180', 'E5EEF3', 'F4F9FC'
SKY, SKY_L, SUN, WHITE, SOFT = '0284C7', '0EA5E9', 'FFCC4D', 'FFFFFF', 'B9CAD6'
FONT = '맑은 고딕'
L, CW = 12.7, 184.7  # 앞면과 같은 좌우 여백·본문 폭(mm)


def set_font(run, size, color, bold=False):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, FONT
    f.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set('typeface', FONT)


def text(slide, x, y, w, h, paras, align='l', anchor='t', name=None, spacing=None):
    """paras: [[(글자, pt, 색, 굵게)], ...] — 문단마다 run 목록"""
    sh = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    if name:
        sh.name = name
    tf = sh.text_frame
    tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        for s, size, color, bold in runs:
            r = p.add_run()
            r.text = s
            set_font(r, size, color, bold)
    return sh


def box(slide, x, y, w, h, fill, line=None, radius=None, shadow=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    sh = slide.shapes.add_shape(shape, Mm(x), Mm(y), Mm(w), Mm(h))
    if name:
        sh.name = name
    style = sh._element.find(qn('p:style'))  # 테마 스타일(그림자·글자색)을 떼어 앞면과 같은 맨 도형으로
    if style is not None:
        sh._element.remove(style)
    if radius is not None:
        sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    if shadow:  # 앞면 카드와 같은 옅은 그림자
        sppr = sh._element.spPr
        sppr.append(etree.fromstring(
            '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:outerShdw blurRad="101600" dist="25400" dir="5400000" algn="bl" rotWithShape="0">'
            '<a:srgbClr val="0C1B27"><a:alpha val="10000"/></a:srgbClr></a:outerShdw></a:effectLst>'))
    return sh


def section_title(slide, x, y, w, title, sub=None, name=None):
    text(slide, x, y, w, 7.6, [[(title, 14, INK, True)]], name=name)
    if sub:
        text(slide, x, y + 1.6, w, 5, [[(sub, 8, MUTED, False)]], align='r')


def build():
    prs = Presentation(SRC)
    front = prs.slides[0]
    blobs = {sh.name: sh for sh in front.shapes if sh.shape_type == 13}
    slide = prs.slides.add_slide(front.slide_layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)

    # ── 머리 띠: 앞면 히어로 배경의 위쪽만 잘라 쓴다 ──
    hero = blobs['hero-bg']
    pic = slide.shapes.add_picture(io.BytesIO(hero.image.blob), 0, 0, prs.slide_width, Mm(30))
    pic.name = 'back-hero-bg'
    pic.crop_bottom = 1 - 30 / (hero.height / 36000)
    logo = blobs['logo']
    slide.shapes.add_picture(io.BytesIO(logo.image.blob), Mm(L), Mm(7.4), logo.width, logo.height).name = 'back-logo'
    text(slide, 26.4, 6.6, 50.8, 9.1, [[('PedalLink', 18, WHITE, True)]], name='back-brand')
    text(slide, L, 17.2, 110, 10, [[('이렇게 ', 20, WHITE, True), ('설계', 20, SUN, True), ('했어요', 20, WHITE, True)]],
         name='back-headline')
    text(slide, 100, 20.3, 97.4, 5, [[('시스템 구성 · 라이딩 흐름 · 데이터 구조 · 설계 포인트', 8.5, 'DCEFFB', True)]],
         align='r', name='back-eyebrow')

    # ── 01 시스템 구성 ──
    section_title(slide, L, 35.5, CW, '시스템 구성', '브라우저부터 데이터베이스까지, 요청이 지나가는 길', 'back-sec1-title')
    box(slide, L, 44, CW, 80, WHITE, LINE, 0.0524 * 0.6, True, name='back-card-arch')
    slide.shapes.add_picture(os.path.join(HERE, 'architecture.png'), Mm(L + 2.5), Mm(46.5), Mm(179.7), Mm(75)).name = 'back-img-arch'

    # ── 02 라이딩 흐름 ──
    section_title(slide, L, 129.5, CW, '라이딩 흐름', '아래 칸은 각 단계의 데이터가 저장되는 곳', 'back-sec2-title')
    box(slide, L, 138, CW, 40, WHITE, LINE, 0.0524 * 1.2, True, name='back-card-flow')
    slide.shapes.add_picture(os.path.join(HERE, 'flow.png'), Mm(L + 2.5), Mm(140.5), Mm(179.7), Mm(35)).name = 'back-img-flow'

    # ── 03 데이터 구조 / 설계 포인트 ──
    section_title(slide, L, 183.5, 118.7, '데이터 구조', '테이블 18개 · PostgreSQL + PostGIS', 'back-sec3-title')
    box(slide, L, 192, 118.7, 79, WHITE, LINE, 0.0524 * 0.6, True, name='back-card-data')
    slide.shapes.add_picture(os.path.join(HERE, 'data.png'), Mm(L + 2.5), Mm(194.5), Mm(113.7), Mm(74)).name = 'back-img-data'

    px, pw = 135.4, 62.0
    text(slide, px, 183.5, pw, 7.6, [[('설계 포인트', 14, INK, True)]], name='back-sec4-title')
    points = [  # 본문은 한국어가 낱자 단위로 줄바꿈되지 않게 두 줄로 미리 나눈다
        ('실시간 위치는 메모리에', ['3초마다 오는 위치는 DB 대신 서버 메모리에', '15초만 두고, 공유를 끄면 바로 지웁니다.']),
        ('거리 계산은 PostGIS로', ['서울 구 경계로 코스 지역을 정하고,', '220m 안 라이더만 매칭·인원 체크합니다.']),
        ('상승고도는 탄 만큼만', ['코스 상승고도 × (탄 거리 ÷ 코스 거리)만', '인정해 통계가 부풀지 않습니다.']),
        ('AI에는 숫자만 보냅니다', ['이름·이메일 없이 통계 숫자만 Claude에 보내고,', '통계가 바뀔 때만 다시 분석합니다.']),
    ]
    ch, cg = 18.25, 2.0
    for i, (title, body) in enumerate(points):
        y = 192 + i * (ch + cg)
        box(slide, px, y, pw, ch, BG2, LINE, 0.12, name=f'back-point-{i + 1}')
        badge = box(slide, px + 3, y + 3, 6.4, 6.4, SKY, shape=MSO_SHAPE.OVAL, name=f'back-point-badge-{i + 1}')
        tf = badge.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f'{i + 1:02d}'
        set_font(r, 7, WHITE, True)
        text(slide, px + 11.6, y + 3.4, pw - 14, 5.6, [[(title, 9.5, INK, True)]], anchor='m', name=f'back-point-title-{i + 1}')
        text(slide, px + 3, y + 10.4, pw - 6, 7.4, [[(line, 7.2, MUTED, False)] for line in body], spacing=1.0, name=f'back-point-body-{i + 1}')

    # ── 바닥 띠 ──
    box(slide, 0, 277, 210.1, 20, INK, shape=MSO_SHAPE.RECTANGLE, name='back-footer-bg')
    text(slide, L, 281.2, 90, 6, [[('PedalLink', 11, WHITE, True)]], name='back-footer-brand')
    text(slide, L, 287.6, 110, 5, [[('서울 자전거 경로 안내 · 협동 라이딩 플랫폼', 7.5, SOFT, False)]], name='back-footer-sub')
    text(slide, 138.4, 281.2, 58.9, 4.6, [[('GitHub', 7.5, SKY_L, True)]], name='back-footer-gh-label')
    gh = text(slide, 138.4, 285.6, 58.9, 4.3, [[('github.com/hosung5287-creator/Pedal_link', 7.5, WHITE, False)]], name='back-footer-gh')
    gh.text_frame.paragraphs[0].runs[0].hyperlink.address = GITHUB
    text(slide, 138.4, 290.2, 58.9, 4, [[('설계 기준 2026. 10 · khs 브랜치 코드', 6.5, SOFT, False)]], name='back-footer-note')

    prs.save(OUT)
    print('saved', os.path.normpath(OUT))


if __name__ == '__main__':
    build()
